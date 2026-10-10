"""House-style output: rebuild the converted paper inside the publisher's own Word file.

Given an original .docx of the same series (e.g. the Word file the PDF was printed
from), the converted content is poured into THAT document's skeleton:

* the cover lines, PW header/logo, footers and every section break (1-column
  title band -> 2-column questions -> ANSWER KEY band -> 2-column key ->
  2-column solutions) are the template's own elements, copied verbatim;
* every paragraph gets the template's exact paragraph properties: hanging indent,
  tab stops (567/992/2552/2977 twips), line spacing, justification;
* questions are ``N.<tab>stem`` with a bold number, options are
  ``<tab>(1)<tab>value<tab>(2)<tab>value`` rows (2-up) or one per row, a blank
  paragraph separates questions, the key is ``N.<tab>(k)`` and every solution
  starts with ``N.<tab>(k)`` followed by ``<tab>text`` lines - as in the source.

Equations stay native Word (OMML) equations; the template's MathType objects are
not reused.
"""
from __future__ import annotations

import copy
import io
import re
from dataclasses import dataclass, field

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROLE = "⟪R:{}⟫"
ROLE_RE = re.compile("^⟪R:(\\w+)⟫")
TAB = "⟪TAB⟫"

QNUM_RE = re.compile(r"^\s*\*?\s*(\d{1,3})\.(?!\d)\s*")  # "12." and starred "*12."
KEY_RE = re.compile(r"^\s*(\d{1,3})\.\s*\(?\s*([1-5a-dA-D])\s*\)?\s*$")
SOL_RE = re.compile(r"^\s*(\d{1,3})\.\s*\(\s*([1-5a-dA-D])\s*\)\s*(.*)$", re.S)
ANSWER_RE = re.compile(r"ANSWER\s*KEY", re.I)
SOLUTION_RE = re.compile(r"\b(SOLUTIONS?|HINTS?)\b", re.I)

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
IMAGE_RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
DROPPABLE_RT = ("/image", "/oleObject", "/package", "/hyperlink", "/chart")

# value-cell widths of the template's tab grid, in points
TWO_UP_CELL = (2552 - 992) / 20 - 4
ONE_UP_START = 992 / 20


SUBJECT_RE = re.compile(r"^(?:PHYSICS|CHEMISTRY|BIOLOGY|BOTANY|ZOOLOGY|MATHEMATICS|MATHS"
                        r"|SECTION[\s\-–]*[A-Z0-9]{1,2}|PART[\s\-–]*[A-Z0-9]{1,2}"
                        r"|[A-Z][A-Z &\-–]{2,39})$")


def is_subject(s: str) -> bool:
    """'CHEMISTRY', 'SECTION-A' - a heading inside the questions, not a question."""
    s = s.strip()
    return bool(s) and not QNUM_RE.match(s) and bool(SUBJECT_RE.match(s))


def question_start(nums: list[int]) -> int:
    """Position in ``nums`` (question numbers in reading order) where the paper starts.

    The paper is the first long run of numbers. A run ends when numbering goes back
    to (near) the run's own start - a restart; small disorder (49, 44) does not end
    it. Numbered cover instructions "1. 2. 3. 4." are a short run and are skipped;
    the answer key and solutions are later runs.
    """
    if not nums:
        return 0
    runs, start = [], 0
    for i in range(1, len(nums) + 1):
        if i == len(nums) or (nums[i] <= nums[i - 1] and nums[i] <= nums[start] + 2):
            runs.append((start, len(set(nums[start:i]))))
            start = i
    longest = max(n for _, n in runs)
    return next(s for s, n in runs if n >= 0.6 * longest)


def _text(p) -> str:
    return "".join(t.text or "" for t in p.iter(qn("w:t")))


def _full_text(p) -> str:
    """Text with tabs, so "1.<tab>stem" can be told from a cover instruction."""
    out = []
    for el in p.iter(qn("w:t"), qn("w:tab")):
        if el.getparent().tag != qn("w:r"):  # tab-stop definitions in w:pPr/w:tabs
            continue
        out.append("\t" if el.tag == qn("w:tab") else (el.text or ""))
    return "".join(out)


def _has_break(p) -> bool:
    pPr = p.find(qn("w:pPr"))
    return pPr is not None and pPr.find(qn("w:sectPr")) is not None


def _is_p(el) -> bool:
    return el.tag == qn("w:p")


# --------------------------------------------------------------------------- #
# A question never breaks across columns or pages
# --------------------------------------------------------------------------- #
_TRPR_AFTER_CANTSPLIT = ("w:trHeight", "w:tblHeader", "w:tblCellSpacing", "w:jc", "w:hidden",
                         "w:ins", "w:del", "w:trPrChange")


def _switch_on(p, tag: str) -> None:
    """w:keepNext / w:keepLines on a paragraph, in schema order (after w:pStyle)."""
    pPr = p.find(qn("w:pPr"))
    if pPr is None:
        pPr = p.makeelement(qn("w:pPr"), {})
        p.insert(0, pPr)
    el = pPr.find(qn(tag))
    if el is not None:
        el.attrib.pop(qn("w:val"), None)
        return
    el = pPr.makeelement(qn(tag), {})
    before = ("w:pStyle", "w:keepNext") if tag == "w:keepLines" else ("w:pStyle",)
    anchor = next((a for a in (pPr.find(qn(t)) for t in reversed(before)) if a is not None), None)
    if anchor is not None:
        anchor.addnext(el)
    else:
        pPr.insert(0, el)


def _is_blank(el) -> bool:
    return _is_p(el) and not _text(el).strip() and not any(
        True for _ in el.iter(qn("w:drawing"), qn("w:object"), qn("w:pict"), qn("m:oMath")))


def keep_together(block: list) -> None:
    """Keep one question (stem, figures, tables, options) in one column of one page.

    Every paragraph keeps its lines together and stays with the next one; table
    rows do not split. The question's last paragraph is free, so the following
    question may start in the next column. A question taller than a whole column
    still breaks - Word then ignores the request.
    """
    block = [el for el in block if _is_p(el) or el.tag == qn("w:tbl")]
    while block and _is_blank(block[-1]):  # the blank line between questions
        block.pop()
    for i, el in enumerate(block):
        last = i == len(block) - 1
        if _is_p(el):
            _switch_on(el, "w:keepLines")
            if not last:
                _switch_on(el, "w:keepNext")
            continue
        rows = el.findall(qn("w:tr"))
        for r, tr in enumerate(rows):
            trPr = tr.find(qn("w:trPr"))
            if trPr is None:
                trPr = tr.makeelement(qn("w:trPr"), {})
                at = next((k for k, c in enumerate(tr) if c.tag != qn("w:tblPrEx")), len(tr))
                tr.insert(at, trPr)
            if trPr.find(qn("w:cantSplit")) is None:
                cs = trPr.makeelement(qn("w:cantSplit"), {})
                later = next((trPr.find(qn(t)) for t in _TRPR_AFTER_CANTSPLIT
                              if trPr.find(qn(t)) is not None), None)
                later.addprevious(cs) if later is not None else trPr.append(cs)
            if last and r == len(rows) - 1:
                continue
            for p in tr.iter(qn("w:p")):  # a table stays with what follows row by row
                _switch_on(p, "w:keepNext")


def lower_option_figures(p) -> None:
    """"(A) [figure]": the picture is lowered by half its height, so the label sits
    beside its middle as printed, not at its foot."""
    if not _text(p).strip():
        return
    for r in p.iter(qn("w:r")):
        ext = next(r.iter(qn("wp:extent")), None)
        if ext is None:
            continue
        drop = int(ext.get("cy")) / 12700 / 2 - 4  # points
        if drop < 6:
            continue
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            rpr = r.makeelement(qn("w:rPr"), {})
            r.insert(0, rpr)
        for old in rpr.findall(qn("w:position")):
            rpr.remove(old)
        pos = rpr.makeelement(qn("w:position"), {qn("w:val"): str(-int(drop * 2))})
        later = next((e for e in rpr if e.tag in _RPR_AFTER_POSITION), None)
        later.addprevious(pos) if later is not None else rpr.append(pos)


# --------------------------------------------------------------------------- #
# Template skeleton
# --------------------------------------------------------------------------- #
@dataclass
class Skeleton:
    doc: object
    preamble: list = field(default_factory=list)     # cover lines + 1-col break
    q_close: object = None                           # para carrying the 2-col break
    answer_head: list = field(default_factory=list)  # "ANSWER KEY" band + 1-col break
    key_close: object = None
    solution_head: list = field(default_factory=list)
    sol_close: object = None
    stem_pPr: object = None
    opt_pPr: object = None
    cont_pPr: object = None
    key_pPr: object = None
    sol_pPr: object = None
    solcont_pPr: object = None
    opt1_pPr: object = None      # long option alone on its row: hangs at 992
    tbl_pr: object = None        # first table's tblPr (style, indent, look)
    cell_pPr: object = None      # a body-row cell paragraph of that table
    head_cell_pPr: object = None  # a header-row cell paragraph (centred)
    per_q: dict = field(default_factory=dict)    # {N: {stem/opt/opt1/cont: pPr}}
    per_sol: dict = field(default_factory=dict)  # {N: {sol/solcont: pPr}}
    subject_bands: dict = field(default_factory=dict)  # {"CHEMISTRY": [elements]}
    q_tail: list = field(default_factory=list)   # after the last question (no key part)
    all_text: str = ""                           # the template's whole text


def _either(a, b):
    return a if a is not None else b


def _per_number(paras: list, head_re: re.Pattern, kinds: dict) -> dict[int, dict]:
    """{N: {kind: pPr}} - each question's / solution's own paragraph settings.

    When the template is the very paper being converted, question N then gets
    exactly the spacing the editor gave question N, so pages break where they did.
    """
    out: dict[int, dict] = {}
    cur = None
    for p in paras:
        m = head_re.match(_text(p))
        if m:
            n = int(m.group(1))
            cur = n if cur is None or n > cur else None
        if cur is None:
            continue
        for kind, test in kinds.items():  # first matching kind wins
            if test(p):
                if kind not in out.setdefault(cur, {}):
                    out[cur][kind] = _pick_pPr([p], want_normal=False)
                break
    return out


def _pick_pPr(paras: list, want_normal: bool = True):
    """First paragraph's pPr, preferring the plain Normal style (11 pt runs)."""
    def style(p):
        ps = p.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
        return ps.get(qn("w:val")) if ps is not None else "Normal"
    pool = [p for p in paras if p.find(qn("w:pPr")) is not None]
    if want_normal:
        pool = [p for p in pool if style(p) == "Normal"] or pool
    if not pool:
        return None
    pPr = copy.deepcopy(pool[0].find(qn("w:pPr")))
    for tag in ("w:sectPr", "w:rPr"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    return pPr


def load_skeleton(path) -> Skeleton:
    doc = Document(str(path))
    body = doc.element.body
    kids = [k for k in body.iterchildren() if k.tag != qn("w:sectPr")]
    sk = Skeleton(doc, all_text=_text(body))

    head_re = re.compile(r"^\s*\*?\s*(\d{1,3})\.\s*\t")
    numbered = [(i, int(m.group(1))) for i, k in enumerate(kids)
                if _is_p(k) and (m := head_re.match(_full_text(k)))]
    # numbered cover instructions ("1.<tab>Immediately fill in...") are not question 1
    first_q = numbered[question_start([n for _, n in numbered])][0]
    sk.preamble = kids[:first_q]

    marker = next((i for i, k in enumerate(kids) if i > first_q and _is_p(k)
                   and ANSWER_RE.search(_text(k))), None)
    q_end = marker if marker is not None else len(kids)
    q_paras = [k for k in kids[first_q:q_end] if _is_p(k)]
    stems = [p for p in q_paras if QNUM_RE.match(_text(p))]
    opts = [p for p in q_paras if re.match(r"^\t?\(\s*(?:[1-4]|[A-Da-d])\s*\)", _full_text(p))]
    conts = [p for p in q_paras if _full_text(p).startswith("\t") and p not in opts]
    sk.per_q = _per_number(q_paras, QNUM_RE, {
        "stem": lambda p: QNUM_RE.match(_text(p)),
        "opt": lambda p: re.match(r"^\t?\(\s*(?:[1-4]|[A-Da-d])\s*\)\t.*\t\(", _full_text(p)),
        "opt1": lambda p: re.match(r"^\t?\(\s*(?:[1-4]|[A-Da-d])\s*\)", _full_text(p)),
        "cont": lambda p: _full_text(p).startswith("\t"),
    })
    sk.stem_pPr = _pick_pPr(stems)
    sk.opt_pPr = _either(_pick_pPr(opts), copy.deepcopy(sk.stem_pPr))
    sk.cont_pPr = _either(_pick_pPr(conts), copy.deepcopy(sk.opt_pPr))
    for jc in sk.opt_pPr.findall(qn("w:jc")):  # option rows are never justified
        sk.opt_pPr.remove(jc)
    hang = [p for p in opts if (lambda i: i is not None and i.get(qn("w:hanging")) == "992")(
        p.find(f"{qn('w:pPr')}/{qn('w:ind')}"))]
    sk.opt1_pPr = _pick_pPr(hang)
    if sk.opt1_pPr is None:  # template has none: derive it from the option rows
        sk.opt1_pPr = copy.deepcopy(sk.opt_pPr)
        ind = sk.opt1_pPr.find(qn("w:ind"))
        if ind is not None:
            ind.set(qn("w:left"), "992")
            ind.set(qn("w:hanging"), "992")

    tbl = next((k for k in kids[first_q:q_end] if k.tag == qn("w:tbl")), None)
    if tbl is not None:
        sk.tbl_pr = copy.deepcopy(tbl.find(qn("w:tblPr")))
        rows = tbl.findall(qn("w:tr"))
        sk.head_cell_pPr = _pick_pPr([p for p in rows[0].iter(qn("w:p"))])
        sk.cell_pPr = _pick_pPr([p for r in rows[1:] for p in r.iter(qn("w:p"))])
        if sk.cell_pPr is None:
            sk.cell_pPr = sk.head_cell_pPr

    # subject bands ("CHEMISTRY" between two 2-column sections): from the blank
    # paragraphs after the previous question to the next question, verbatim
    for h in range(first_q, q_end):
        k = kids[h]
        if not (_is_p(k) and is_subject(_text(k))):
            continue
        pPr = k.find(qn("w:pPr"))
        structural = _has_break(k) or (pPr is not None and pPr.find(qn("w:shd")) is not None) \
            or (h > 0 and _is_p(kids[h - 1]) and _has_break(kids[h - 1]))
        if not structural:
            continue
        a = h
        while a - 1 > first_q and _is_p(kids[a - 1]) and not _text(kids[a - 1]).strip() \
                and kids[a - 1].find(".//" + qn("w:drawing")) is None:
            a -= 1
        b = next((i for i in range(h + 1, q_end) if _is_p(kids[i])
                  and head_re.match(_full_text(kids[i]))), q_end)
        sk.subject_bands.setdefault(re.sub(r"[^A-Z0-9]", "", _text(k).upper()), kids[a:b])

    # the section break after the LAST question closes the 2-column part (later
    # breaks belong to the closing ornament / answer-key band)
    last_q = max((i for i, _ in numbered if i < q_end), default=first_q)
    breaks = [i for i in range(last_q, q_end) if _is_p(kids[i]) and _has_break(kids[i])]
    if not breaks:
        breaks = [i for i in range(first_q, q_end) if _is_p(kids[i]) and _has_break(kids[i])][-1:]
    q_close_i = breaks[0] if breaks else None
    sk.q_close = kids[q_close_i] if q_close_i is not None else None
    if marker is None:
        if q_close_i is not None:
            sk.q_tail = kids[q_close_i + 1:]
        return sk

    first_key = next(i for i in range(marker, len(kids))
                     if _is_p(kids[i]) and KEY_RE.match(_text(kids[i])))
    sk.answer_head = kids[(q_close_i + 1 if q_close_i is not None else marker):first_key]
    key_close_i = next((i for i in range(first_key, len(kids))
                        if _is_p(kids[i]) and _has_break(kids[i])), None)
    keys = [k for k in kids[first_key:key_close_i or len(kids)]
            if _is_p(k) and KEY_RE.match(_text(k))]
    sk.key_pPr = _pick_pPr(keys)
    if key_close_i is None:
        return sk
    sk.key_close = kids[key_close_i]

    first_sol = next((i for i in range(key_close_i + 1, len(kids))
                      if _is_p(kids[i]) and SOL_RE.match(_text(kids[i]))), None)
    if first_sol is None:
        return sk
    sk.solution_head = kids[key_close_i + 1:first_sol]
    rest = kids[first_sol:]
    sol_breaks = [k for k in rest if _is_p(k) and _has_break(k)]
    sk.sol_close = sol_breaks[-1] if sol_breaks else None
    starts = [k for k in rest if _is_p(k) and SOL_RE.match(_text(k))]
    sconts = [k for k in rest if _is_p(k) and _full_text(k).startswith("\t")
              and not SOL_RE.match(_text(k))]
    sk.per_sol = _per_number([k for k in rest if _is_p(k)], SOL_RE, {
        "sol": lambda p: SOL_RE.match(_text(p)),
        "solcont": lambda p: _full_text(p).startswith("\t"),
    })
    sk.sol_pPr = _pick_pPr(starts)
    sk.solcont_pPr = _either(_pick_pPr(sconts), copy.deepcopy(sk.sol_pPr))
    return sk


# --------------------------------------------------------------------------- #
# Roles: split the converted lines into cover / questions / key / solutions
# --------------------------------------------------------------------------- #
def split_parts(lines: list[str]) -> dict[str, list[str]]:
    parts = {"pre": [], "q": [], "key": [], "sol": []}
    kept, prev = [], None
    for ln in lines:
        s = re.sub(r"<[^>]+>", "", ln).strip()
        if not s and not ln.strip():
            continue
        if s and s == prev:  # the same line twice (layout block + recovered text)
            continue
        prev = s
        kept.append((ln, s))
    # where the questions begin: not at numbered cover instructions
    numbered = [(i, int(m.group(1))) for i, (_, s) in enumerate(kept) if (m := QNUM_RE.match(s))]
    begin = numbered[question_start([n for _, n in numbered])][0] if numbered else len(kept)
    phase, first_q, last_q, last_key = "pre", 0, 0, 0
    for i, (ln, s) in enumerate(kept):
        if ln.lstrip().startswith("<h2>"):
            if ANSWER_RE.search(s) and len(s) < 40:
                phase = "key"
                continue
            if SOLUTION_RE.search(s) and len(s) < 40 and phase in ("key", "q"):
                phase = "sol"
                continue
            if QNUM_RE.match(s):  # a question the layout model took for a heading
                ln = s
        m = QNUM_RE.match(s)
        if phase == "pre" and m and i >= begin:
            phase, first_q = "q", int(m.group(1))
        elif phase == "q" and m and int(m.group(1)) <= first_q and last_q > first_q:
            phase = "key" if KEY_RE.match(s) else "sol"   # numbering restarted
        if phase == "key":
            k = KEY_RE.match(s)
            if k and int(k.group(1)) > last_key:
                last_key = int(k.group(1))
            elif k or SOL_RE.match(s):
                phase = "sol"                              # key done, solutions begin
            elif not s:
                continue
        if phase == "q" and m:
            last_q = int(m.group(1))
        parts[phase].append(ln)
    return parts


def option_rows(opts: list[str], per: int) -> list[str]:
    """['(1) a', '(2) b'] -> ['<TAB>(1)<TAB>a<TAB>(2)<TAB>b'] (per=2)."""
    rows = []
    for k in range(0, len(opts), per):
        cells = []
        for o in opts[k:k + per]:
            m = re.match(r"\s*(\([^)]{1,3}\))\s*(.*)$", o, re.S)
            cells.append(TAB + m.group(1) + TAB + m.group(2).strip() if m else TAB + o)
        rows.append("".join(cells))
    return rows


def tag_roles(parts: dict[str, list[str]], is_option_row) -> list[str]:
    """Prefix every paragraph with its role token (tables are left as they are)."""
    out = []
    topic = next((ln for ln in parts["pre"] if re.sub(r"<[^>]+>", "", ln).strip()), None)
    if topic:  # plain text: a "<h2>" here would become a paragraph without its role
        out.append(ROLE.format("topic") + re.sub(r"<[^>]+>", "", topic).strip())
    first = True
    for ln in parts["q"]:
        if ln.lstrip().startswith("<table"):
            out.append(ln)
            continue
        s = re.sub(r"<[^>]+>", "", ln).strip()
        if is_subject(s) and (ln.lstrip().startswith("<h2>") or re.fullmatch(
                r"PHYSICS|CHEMISTRY|BIOLOGY|BOTANY|ZOOLOGY|MATHEMATICS", s)):
            out.append(ROLE.format("subject") + s)  # "CHEMISTRY" band between subjects
            first = True
            continue
        if QNUM_RE.match(s):
            if not first:
                out.append(ROLE.format("blank"))
            first = False
            out.append(ROLE.format("stem") + ln)
        elif is_option_row(ln):  # one option per row hangs at the value tab
            out.append(ROLE.format("opt1" if ln.count(TAB) <= 2 else "opt") + ln)
        else:  # "I. Graph..." / "A. For Person A..." -> <tab>I.<tab>Graph...
            ln = re.sub(r"^\s*(VI|IV|V|I{1,3}|[A-E])\.\s+", lambda m: m.group(1) + "." + TAB, ln)
            ln = re.sub(r"^\s*([•●▪■◦➢➤►])\s*", lambda m: m.group(1) + TAB, ln)  # <tab>•<tab>item
            if re.fullmatch(r"[A-E](?: [A-E]){1,4}", s):  # "A B C D" over a T/F grid
                ln = TAB + TAB.join(s.split())
            out.append(ROLE.format("cont") + ln)
    for ln in parts["key"]:
        out.append(ROLE.format("key") + ln)
    first = True
    for ln in parts["sol"]:
        if ln.lstrip().startswith("<table"):
            out.append(ln)
            continue
        s = re.sub(r"<[^>]+>", "", ln).strip()
        m = SOL_RE.match(s)
        if m:
            if not first:
                out.append(ROLE.format("blank_sol"))
            first = False
            out.append(ROLE.format("sol") + f"{m.group(1)}. ({m.group(2)})")
            if m.group(3).strip():
                rest = re.sub(r"^.*?\(\s*[1-5a-dA-D]\s*\)\s*", "", ln, count=1)
                out.append(ROLE.format("solcont") + rest)
        else:
            out.append(ROLE.format("solcont") + ln)
    return out


# --------------------------------------------------------------------------- #
# Build: pour the Pandoc paragraphs into the template skeleton
# --------------------------------------------------------------------------- #
def _new_p(pPr):
    el = OxmlElement("w:p")
    if pPr is not None:
        el.append(copy.deepcopy(pPr))
    return el


def _run(text: str, bold: bool = False, size: int = 22):
    r = OxmlElement("w:r")
    rpr = r.makeelement(qn("w:rPr"), {})
    if bold:
        rpr.append(rpr.makeelement(qn("w:b"), {}))
    for tag in ("w:sz", "w:szCs"):
        e = rpr.makeelement(qn(tag), {qn("w:val"): str(size)})
        rpr.append(e)
    r.append(rpr)
    parts = text.split("\t")
    for i, piece in enumerate(parts):
        if i:
            r.append(r.makeelement(qn("w:tab"), {}))
        if piece:
            t = r.makeelement(qn("w:t"), {"{http://www.w3.org/XML/1998/namespace}space": "preserve"})
            t.text = piece
            r.append(t)
    return r


def _break_only(p):
    """Copy of a template paragraph carrying a section break, without its own text
    (the template's last option/key line would otherwise reappear in the output)."""
    el = copy.deepcopy(p)
    for c in _content_children(el):
        el.remove(c)
    return el


def _content_children(p):
    return [c for c in p if c.tag != qn("w:pPr") and c.tag not in
            (qn("w:bookmarkStart"), qn("w:bookmarkEnd"))]


def _strip_leading(children: list, pattern: re.Pattern) -> re.Match | None:
    """Remove the regex match from the start of the paragraph's leading text runs."""
    text, runs = "", []
    for c in children:
        if c.tag != qn("w:r"):
            break
        runs.append(c)
        text += "".join(t.text or "" for t in c.iter(qn("w:t")))
    m = pattern.match(text)
    if not m:
        return None
    left = m.end()
    for r in runs:
        for t in r.iter(qn("w:t")):
            n = len(t.text or "")
            take = min(n, left)
            t.text = (t.text or "")[take:]
            left -= take
        if left <= 0:
            break
    return m


_RPR_AFTER_SZ = {qn("w:" + t) for t in (
    "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em",
    "lang", "eastAsianLayout", "specVanish", "oMath")}
_RPR_AFTER_POSITION = _RPR_AFTER_SZ | {qn("w:sz"), qn("w:szCs")}
_PPR_AFTER_JC = {qn("w:" + t) for t in (
    "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle",
    "rPr", "sectPr", "pPrChange")}


def _normalise_runs(children: list, size: int = 22) -> None:
    for c in children:
        for r in ([c] if c.tag == qn("w:r") else c.iter(qn("w:r"))):
            rpr = r.find(qn("w:rPr"))
            if rpr is None:
                rpr = r.makeelement(qn("w:rPr"), {})
                r.insert(0, rpr)
            for tag in ("w:rStyle", "w:rFonts", "w:sz", "w:szCs"):
                for e in rpr.findall(qn(tag)):
                    rpr.remove(e)
            if r.tag == qn("w:r") and c.tag == qn("w:r"):
                # schema order: w:sz/w:szCs come before highlight, u, vertAlign, lang ...
                later = next((e for e in rpr if e.tag in _RPR_AFTER_SZ), None)
                for tag in ("w:sz", "w:szCs"):
                    e = rpr.makeelement(qn(tag), {qn("w:val"): str(size)})
                    later.addprevious(e) if later is not None else rpr.append(e)
            for t in r.iter(qn("w:t")):  # tab placeholders -> real tabs
                if t.text and TAB in t.text:
                    pieces = t.text.split(TAB)
                    t.text = pieces[0]
                    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                    anchor = t
                    for piece in pieces[1:]:
                        tab = r.makeelement(qn("w:tab"), {})
                        anchor.addnext(tab)
                        nt = r.makeelement(qn("w:t"), {"{http://www.w3.org/XML/1998/namespace}space": "preserve"})
                        nt.text = piece
                        tab.addnext(nt)
                        anchor = nt


def _display_to_inline(children: list) -> list:
    """m:oMathPara -> its m:oMath children (display maths sit after a tab, inline)."""
    out = []
    for c in children:
        if c.tag == qn("m:oMathPara"):
            out += [m for m in c if m.tag == qn("m:oMath")]
        else:
            out.append(c)
    return out


def build(sk: Skeleton, pandoc_docx, out_path, fit_table=None, col_width_pt: float = 255,
          centred: frozenset = frozenset()) -> None:
    """``centred``: file names of the figures the PDF prints in the middle of a column."""
    src = Document(str(pandoc_docx))
    doc = sk.doc
    body = doc.element.body
    final_sectPr = body.find(qn("w:sectPr"))
    for k in list(body.iterchildren()):
        if k is not final_sectPr:
            body.remove(k)

    blocks: list[list] = [[]]  # body elements per question; blocks[-1] is the open one
    in_question = False

    def add(el):
        final_sectPr.addprevious(el)
        if in_question:
            blocks[-1].append(el)

    # images: copy the Pandoc doc's pictures into the template package
    img_map = {}

    def remap_images(el):
        for blip in el.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
            rid = blip.get(f"{{{R_NS}}}embed")
            if rid and rid not in img_map:
                blob = src.part.related_parts[rid].blob
                img_map[rid] = doc.part.get_or_add_image(io.BytesIO(blob))[0]
            if rid:
                blip.set(f"{{{R_NS}}}embed", img_map[rid])
        for i, dp in enumerate(el.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr")):
            dp.set("id", str(5000 + len(img_map) * 10 + i))

    # 1. cover: template's own lines; the topic line text is replaced
    items = list(src.element.body.iterchildren())
    roles = []
    for el in items:
        if el.tag == qn("w:p"):
            kids = _content_children(el)
            m = _strip_leading(kids, ROLE_RE)
            roles.append((m.group(1) if m else "cont", el, kids))
        elif el.tag == qn("w:tbl"):
            roles.append(("table", el, None))

    topic = next((k for r, _, k in roles if r == "topic"), None)
    pre = [copy.deepcopy(e) for e in sk.preamble]
    norm = lambda s: re.sub(r"[^a-z0-9]", "", s.lower())  # noqa: E731
    topic_text = "".join(t.text or "" for c in (topic or []) for t in c.iter(qn("w:t")))
    if topic is not None and norm(topic_text) and norm(topic_text) in norm(sk.all_text):
        topic = None  # the template already says it (same paper): keep its cover as is
    if topic is not None:
        tp = next((e for e in pre if _is_p(e) and _text(e).strip()), None)
        if tp is not None:
            new_text = "".join(t.text or "" for c in topic for t in c.iter(qn("w:t"))).strip()
            ts = list(tp.iter(qn("w:t")))
            if ts and new_text:
                ts[0].text = new_text
                for t in ts[1:]:
                    t.text = ""
    for e in pre:
        add(e)

    def para(pPr, kids, lead: list | None = None):
        p = _new_p(pPr)
        for r in lead or []:
            p.append(r)
        kids = _display_to_inline(kids)
        _normalise_runs(kids)
        for c in kids:
            remap_images(c)
            p.append(c)
        add(p)
        return p

    def own(part: str, n, kind: str, default, nojc: bool = False):
        """Question/solution N's own paragraph settings from the template, if it has N."""
        table = sk.per_q if part == "q" else sk.per_sol
        pPr = table.get(n, {}).get(kind)
        if pPr is None or (kind == "opt1" and pPr is table.get(n, {}).get("opt")):
            return default
        if nojc:
            pPr = copy.deepcopy(pPr)
            for jc in pPr.findall(qn("w:jc")):
                pPr.remove(jc)
        return pPr

    def lone_centred_figure(kids) -> bool:
        names = [c.get("descr", "").rsplit("/", 1)[-1]
                 for k in kids for c in k.iter(qn("pic:cNvPr"))]
        text = "".join(t.text or "" for k in kids for t in k.iter(qn("w:t"))).strip()
        return len(names) == 1 and not text and names[0] in centred

    stage, qn_, sn = "q", None, None
    for role, el, kids in roles:
        if role == "topic":
            continue
        if role == "stem":
            blocks.append([])
            in_question = True
        elif role not in ("opt", "opt1", "cont", "blank", "table"):
            in_question = False  # subject band, answer key, solutions
        if role == "table":
            t = copy.deepcopy(el)
            remap_images(t)
            add(t)
            house_table(sk, t, fit_table, col_width_pt)
            continue
        if role in ("key", "sol", "blank_sol", "solcont") and stage == "q":
            # close the questions' 2-column section, then the template's key band
            add(_break_only(sk.q_close) if sk.q_close is not None else _new_p(sk.stem_pPr))
            for e in sk.answer_head:
                add(copy.deepcopy(e))
            stage = "key"
        if role in ("sol", "blank_sol", "solcont") and stage == "key":
            if sk.key_close is not None:
                add(_break_only(sk.key_close))
            for e in sk.solution_head:
                add(copy.deepcopy(e))
            stage = "sol"

        if role == "stem":
            m = _strip_leading(kids, QNUM_RE)
            qn_ = int(m.group(1)) if m else None
            star = "*" if m and m.group(0).lstrip().startswith("*") else ""
            num = star + m.group(1) + "." if m else ""
            para(own("q", qn_, "stem", sk.stem_pPr), kids, [_run(num, True), _run("\t", True)])
        elif role == "opt":
            lower_option_figures(para(own("q", qn_, "opt", sk.opt_pPr, nojc=True), kids))
        elif role == "opt1":
            lower_option_figures(
                para(own("q", qn_, "opt1", _either(sk.opt1_pPr, sk.opt_pPr), nojc=True), kids))
        elif role == "cont" and lone_centred_figure(kids):
            p = para(own("q", qn_, "cont", sk.cont_pPr), kids)  # in the middle, as printed
            pPr = p.find(qn("w:pPr"))
            if pPr is None:
                pPr = p.makeelement(qn("w:pPr"), {})
                p.insert(0, pPr)
            for old in pPr.findall(qn("w:jc")) + pPr.findall(qn("w:ind")):
                pPr.remove(old)
            jc = pPr.makeelement(qn("w:jc"), {qn("w:val"): "center"})
            later = next((e for e in pPr if e.tag in _PPR_AFTER_JC), None)
            later.addprevious(jc) if later is not None else pPr.append(jc)
        elif role == "cont":
            para(own("q", qn_, "cont", sk.cont_pPr), kids, [_run("\t")])
        elif role == "blank":
            add(_new_p(own("q", qn_, "opt", sk.opt_pPr, nojc=True)))
        elif role == "subject":
            name = "".join(t.text or "" for c in kids for t in c.iter(qn("w:t"))).strip()
            band = sk.subject_bands.get(re.sub(r"[^A-Z0-9]", "", name.upper()))
            if band is None and sk.subject_bands:  # another subject's band, renamed
                band = next(iter(sk.subject_bands.values()))
            if band is None:  # template has no subject bands: a centred bold line
                p = _new_p(sk.stem_pPr)
                for jc in p.findall(f"{qn('w:pPr')}/{qn('w:jc')}"):
                    jc.getparent().remove(jc)
                p.append(_run(name, True))
                add(p)
                continue
            for e in band:
                e = copy.deepcopy(e)
                if _is_p(e) and is_subject(_text(e)):
                    ts = list(e.iter(qn("w:t")))
                    ts[0].text = name
                    for t in ts[1:]:
                        t.text = ""
                add(e)
        elif role == "key":
            k = _strip_leading(kids, re.compile(r"\s*(\d{1,3})\.\s*\(?\s*([1-5a-dA-D])\s*\)?"))
            if k:
                para(_either(sk.key_pPr, sk.stem_pPr), [], [_run(f"{k.group(1)}.", True),
                                                             _run(f"\t({k.group(2)})", True)])
        elif role == "sol":
            k = _strip_leading(kids, re.compile(r"\s*(\d{1,3})\.\s*\(\s*([1-5a-dA-D])\s*\)"))
            if k:
                sn = int(k.group(1))
                para(own("sol", sn, "sol", _either(sk.sol_pPr, sk.stem_pPr)), [],
                     [_run(f"{k.group(1)}.", True), _run(f"\t({k.group(2)})", True)])
        elif role == "blank_sol":
            add(_new_p(own("sol", sn, "solcont", _either(sk.solcont_pPr, sk.cont_pPr))))
        elif role == "solcont":
            para(own("sol", sn, "solcont", _either(sk.solcont_pPr, sk.cont_pPr)), kids,
                 [_run("\t")])

    closing = {"q": sk.q_close, "key": sk.key_close, "sol": sk.sol_close}[stage]
    if closing is not None:
        add(_break_only(closing))
    if stage == "q":  # the template's own ending (end ornament, closing sections)
        for e in sk.q_tail:
            add(copy.deepcopy(e))

    for block in blocks:
        keep_together(block)
    _drop_unused_rels(doc)
    doc.save(str(out_path))


def house_table(sk: Skeleton, t, fit_table, col_width_pt: float) -> None:
    """Template table look: its style/indent/look, its cell paragraphs, 11 pt runs."""
    ind = 28.35
    if sk.tbl_pr is not None:
        tin = sk.tbl_pr.find(qn("w:tblInd"))
        if tin is not None and (tin.get(qn("w:w")) or "").isdigit():
            ind = int(tin.get(qn("w:w"))) / 20
    if fit_table:
        fit_table(t, col_width_pt - ind, ind)
    pr = t.find(qn("w:tblPr"))
    if sk.tbl_pr is not None:
        new = copy.deepcopy(sk.tbl_pr)
        # keep the fitted width and fixed grid, in schema order (Word is strict)
        order = {"w:tblW": ("w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
                            "w:tblLayout", "w:tblCellMar", "w:tblLook"),
                 "w:tblLayout": ("w:tblCellMar", "w:tblLook")}
        for tag, before in order.items():
            old = pr.find(qn(tag))
            for e in new.findall(qn(tag)):
                new.remove(e)
            if old is None:
                continue
            anchor = next((new.find(qn(b)) for b in before if new.find(qn(b)) is not None), None)
            if anchor is not None:
                anchor.addprevious(copy.deepcopy(old))
            else:
                new.append(copy.deepcopy(old))
        pr.addprevious(new)
        t.remove(pr)
    for r, tr in enumerate(t.findall(qn("w:tr"))):
        cell_pPr = _either(sk.head_cell_pPr, sk.cell_pPr) if r == 0 else sk.cell_pPr
        for p in tr.iter(qn("w:p")):
            old = p.find(qn("w:pPr"))
            if old is not None:
                p.remove(old)
            if cell_pPr is not None:
                p.insert(0, copy.deepcopy(cell_pPr))
            _normalise_runs([c for c in p if c.tag != qn("w:pPr")])


def _drop_unused_rels(doc) -> None:
    """The template's 400 MathType objects and pictures are no longer referenced."""
    xml = doc.element.xml
    used = set(re.findall(r'r:(?:id|embed|link|pict)="([^"]+)"', xml))
    rels = doc.part.rels
    store = getattr(rels, "_rels", rels)  # python-docx keeps them in a private dict
    for rid in [r for r, rel in list(rels.items())
                if rel.reltype.endswith(DROPPABLE_RT) and r not in used]:
        store.pop(rid)
